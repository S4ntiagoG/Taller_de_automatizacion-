import openpyxl
import os
import re
import logging
import unicodedata
from datetime import datetime, date
from urllib.parse import quote_plus
from openpyxl.worksheet.table import Table, TableStyleInfo
from dotenv import load_dotenv
from playwright.sync_api import (
    expect,
    sync_playwright,
    TimeoutError as PlaywrightTimeoutError
)

# Cargar configuración
load_dotenv()

URL_SICOF = os.getenv("SICOF_URL", "http://127.0.0.1:8020")
USUARIO = os.getenv("SICOF_USUARIO")
CLAVE = os.getenv("SICOF_CLAVE")

ruta_entrada = os.getenv(
    "ARCHIVO_ENTRADA", "datos/Santiago_Giraldo_por_radicar.xlsx"
)
ruta_salida = os.getenv("ARCHIVO_SALIDA", "salida/resultados.xlsx")
RUTA_LOG = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "salida", "bot.log"
)

MAX_INTENTOS = 3


def configurar_log():
    directorio_salida = os.path.dirname(RUTA_LOG)
    os.makedirs(directorio_salida, exist_ok=True)
    os.makedirs(os.path.join(directorio_salida, "evidencias"), exist_ok=True)

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[
            logging.FileHandler(RUTA_LOG, encoding="utf-8"),
            logging.StreamHandler()
        ],
        force=True
    )


def normalizar_encabezado(texto):
    texto = str(texto or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", texto)


def limpiar_nit(nit):
    if nit is None:
        return ""
    return re.sub(r"\D", "", str(nit))


def normalizar_factura(factura):
    return re.sub(r"[^a-zA-Z0-9]", "", str(factura or "")).upper()


def convertir_fecha(fecha):
    if isinstance(fecha, (datetime, date)):
        return fecha.strftime("%d/%m/%Y")

    texto = str(fecha or "").strip()

    if texto.endswith("Z"):
        try:
            fecha_iso = datetime.fromisoformat(
                texto[:-1] + "+00:00"
            )
            return fecha_iso.strftime("%d/%m/%Y")
        except ValueError:
            pass

    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y",
                    "%Y/%m/%d", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(texto[:19], formato).strftime("%d/%m/%Y")
        except ValueError:
            pass

    raise ValueError("La fecha no tiene un formato válido.")


def validar_valor(valor):
    if valor is None:
        raise ValueError("El valor de la factura está vacío.")

    if isinstance(valor, (int, float)):
        numero = float(valor)
    else:
        texto = str(valor).strip().replace("$", "").replace(" ", "")

        if "," in texto and "." in texto:
            # Formato colombiano: 1.234,56
            texto = texto.replace(".", "").replace(",", ".")
        elif "," in texto:
            texto = texto.replace(",", ".")
        elif texto.count(".") > 1:
            texto = texto.replace(".", "")

        try:
            numero = float(texto)
        except ValueError:
            raise ValueError("El valor de la factura no es numérico.")

    if numero <= 0 or not numero.is_integer():
        raise ValueError(
            "El valor de la factura debe ser un entero mayor que cero."
        )

    return int(numero)


def obtener_datos_fila(fila, columnas):
    # Mantiene compatibilidad con el orden de columnas de tu Excel actual.
    def obtener(nombre, posicion):
        indice = columnas.get(nombre)
        if indice is not None and indice < len(fila):
            return fila[indice]
        return fila[posicion] if posicion < len(fila) else None

    return {
        "id_lista": obtener("idlista", 0),
        "factura": obtener("factura", 1),
        "nit": obtener("nit", 2),
        "proveedor": obtener("proveedor", 3),
        "valor": obtener("valor", 4),
        "centro_costo": obtener("centrocosto", 5),
        "fecha": obtener("fechafactura", 6)
    }


def buscar_causacion(page, nit, factura):
    """Busca una causación que coincida tanto en NIT como en factura."""
    url = f"{URL_SICOF}/causaciones?factura={quote_plus(str(factura))}"
    page.goto(url)
    page.wait_for_load_state("networkidle")

    tabla = page.locator("table.datos tbody")
    filas = tabla.locator("tr")

    for i in range(filas.count()):
        fila = filas.nth(i)
        celdas = fila.locator("td")
        valores = [
            celdas.nth(j).inner_text().strip()
            for j in range(celdas.count())
        ]

        nit_encontrado = any(
            limpiar_nit(valor) == limpiar_nit(nit)
            for valor in valores
            if limpiar_nit(valor)
        )

        factura_encontrada = any(
            normalizar_factura(valor) == normalizar_factura(factura)
            for valor in valores
            if normalizar_factura(valor)
        )

        if nit_encontrado and factura_encontrada:
            codigo = valores[0] if valores else ""
            return True, codigo

    return False, ""


def guardar_resultados(wb_salida, hoja_salida):
    if hoja_salida.max_row > 1:
        referencia = f"A1:F{hoja_salida.max_row}"
        if "Resultados" in hoja_salida.tables:
            hoja_salida.tables["Resultados"].ref = referencia
        else:
            tabla = Table(
                displayName="Resultados",
                ref=referencia
            )
            estilo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            tabla.tableStyleInfo = estilo
            hoja_salida.add_table(tabla)

    wb_salida.save(ruta_salida)


def preparar_y_procesar():
    configurar_log()

    if not USUARIO or not CLAVE:
        raise ValueError(
            "Configura SICOF_USUARIO y SICOF_CLAVE en el archivo .env."
        )

    if not os.path.exists(ruta_entrada):
        raise FileNotFoundError(f"No se encontró el archivo: {ruta_entrada}")

    wb_entrada = openpyxl.load_workbook(ruta_entrada, data_only=True)
    hoja_entrada = wb_entrada["PorRadicar"]

    # Leer columnas por encabezado
    encabezados = {
        normalizar_encabezado(celda.value): indice
        for indice, celda in enumerate(hoja_entrada[1])
        if celda.value is not None
    }

    columnas = {
        "idlista": encabezados.get("idlista"),
        "factura": encabezados.get("factura"),
        "nit": encabezados.get("nit"),
        "proveedor": encabezados.get("proveedor"),
        "valor": encabezados.get("valor"),
        "centrocosto": encabezados.get("centrocosto"),
        "fechafactura": encabezados.get("fechafactura")
    }

    wb_salida = openpyxl.Workbook()
    hoja_salida = wb_salida.active
    hoja_salida.title = "Resultados"
    hoja_salida.append([
        "IdLista", "Factura", "NIT", "Estado", "Causacion", "Detalle"
    ])

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page = browser.new_page()

        try:
            logging.info("Iniciando sesión en SICOF.")
            print("Iniciando sesión en SICOF...")

            page.goto(f"{URL_SICOF}/login")
            page.get_by_label("Usuario").fill(USUARIO)
            page.get_by_label("Contraseña").fill(CLAVE)
            page.get_by_role("button", name="Ingresar").click()
            page.wait_for_load_state("networkidle")

            if page.get_by_label("Usuario").is_visible():
                raise RuntimeError("No fue posible iniciar sesión en SICOF.")

            print("Sesión iniciada. Validando facturas...")

            for numero_fila, fila in enumerate(
                hoja_entrada.iter_rows(min_row=2, values_only=True),
                start=2
            ):
                datos = obtener_datos_fila(fila, columnas)

                id_lista = datos["id_lista"]
                factura = datos["factura"]
                nit = datos["nit"]
                valor = datos["valor"]
                centro_costo = datos["centro_costo"]
                fecha_factura = datos["fecha"]

                if id_lista is None:
                    continue

                estado = ""
                codigo_causacion = ""
                detalle = ""

                try:
                    # Validar antes de navegar por el formulario de SICOF.
                    if not limpiar_nit(nit):
                        raise ValueError("El NIT está vacío o no es válido.")

                    if not str(factura or "").strip():
                        raise ValueError("El número de factura está vacío.")

                    valor_validado = validar_valor(valor)
                    fecha_formateada = convertir_fecha(fecha_factura)

                    if centro_costo is None or not str(centro_costo).strip():
                        raise ValueError("El centro de costo está vacío.")

                    nit_limpio = limpiar_nit(nit)

                    logging.info(
                        "Procesando fila %s: factura %s.",
                        numero_fila, factura
                    )
                    print(f" -> Consultando factura {factura}...")

                    # Idempotencia: NIT y número de factura deben coincidir.
                    existe, codigo = buscar_causacion(
                        page, nit_limpio, factura
                    )

                    if existe:
                        estado = "YA_CAUSADA"
                        codigo_causacion = codigo
                        detalle = "La factura ya estaba registrada."
                        logging.info(
                            "Factura %s ya estaba registrada: %s",
                            factura, codigo
                        )
                        print(f"    Factura ya registrada: {codigo}")

                    else:
                        logging.info(
                            "Factura %s no encontrada; creando causación.",
                            factura
                        )
                        print("    Factura no encontrada. Creando causación...")

                        exito = False

                        for intento in range(1, MAX_INTENTOS + 1):
                            try:
                                
                                page.goto(f"{URL_SICOF}/causaciones/nueva")
                                page.wait_for_load_state("networkidle")

                                expect(
                                    page.get_by_label("NIT del proveedor")
                                ).to_be_editable()

                                page.get_by_label(
                                    "NIT del proveedor"
                                ).fill(nit_limpio)
                                page.get_by_label(
                                    "Número de factura"
                                ).fill(str(factura))
                                page.get_by_label(
                                    "Valor total"
                                ).fill(str(valor_validado))
                                page.get_by_label(
                                    "Fecha de factura"
                                ).fill(fecha_formateada)
                                page.get_by_label(
                                    "Centro de costo"
                                ).select_option(str(centro_costo))

                                page.get_by_role(
                                    "button", name="Continuar"
                                ).click()
                                page.wait_for_load_state("networkidle")

                                error_formulario = page.locator(
                                    ".mensaje-error"
                                )
                                alerta = page.locator('[role="alert"]')

                                if error_formulario.is_visible():
                                    detalle = error_formulario.inner_text()
                                    estado = "RECHAZADA"
                                    break

                                if alerta.is_visible():
                                    detalle = alerta.inner_text()
                                    estado = "RECHAZADA"
                                    break

                                page.get_by_role(
                                    "button", name="Confirmar causación"
                                ).click()
                                page.wait_for_load_state("networkidle")

                                texto_final = page.locator("body").inner_text()

                                match = re.search(
                                    r"CXP-\d{4}-\d+", texto_final
                                )

                                if match:
                                    codigo_causacion = match.group(0)
                                    estado = "CAUSADA"
                                    detalle = "Causación creada correctamente."
                                    exito = True
                                    logging.info(
                                        "Factura %s causada: %s",
                                        factura, codigo_causacion
                                    )
                                    print(
                                        f"    Causación creada: "
                                        f"{codigo_causacion}"
                                    )
                                    break

                                if (
                                    "SICOF-503" in texto_final
                                    or "Servicio temporalmente no disponible"
                                    in texto_final
                                ):
                                    detalle = "SICOF presentó un error temporal."
                                else:
                                    detalle = (
                                        "No se encontró el código de causación "
                                        "en la respuesta de SICOF."
                                    )

                                # Antes de reintentar, comprobar si SICOF
                                # alcanzó a registrar la factura.
                                existe, codigo = buscar_causacion(
                                    page, nit_limpio, factura
                                )

                                if existe:
                                    estado = "YA_CAUSADA"
                                    codigo_causacion = codigo
                                    detalle = (
                                        "La factura quedó registrada en SICOF "
                                        "después de una respuesta incierta."
                                    )
                                    exito = True
                                    break

                                logging.warning(
                                    "Intento %s/%s fallido para factura %s: %s",
                                    intento, MAX_INTENTOS, factura, detalle
                                )

                            except PlaywrightTimeoutError as error:
                                detalle = f"Tiempo de espera agotado: {error}"
                                logging.warning(
                                    "Timeout en factura %s, intento %s: %s",
                                    factura, intento, error
                                )

                            except Exception as error:
                                detalle = str(error)
                                logging.exception(
                                    "Error en factura %s, intento %s",
                                    factura, intento
                                )

                            if intento < MAX_INTENTOS:
                                page.wait_for_timeout(2000)

                        if not exito and estado != "RECHAZADA":
                            estado = "FALLA_SISTEMA"
                            detalle = (
                                f"Se agotaron los {MAX_INTENTOS} intentos. "
                                f"Último detalle: {detalle}"
                            )

                            nombre_captura = re.sub(
                                r"[^a-zA-Z0-9_-]", "_", str(factura)
                            )
                            ruta_captura = (
                                f"salida/evidencias/"
                                f"factura_{nombre_captura}.png"
                            )

                            try:
                                page.screenshot(
                                    path=ruta_captura,
                                    full_page=True
                                )
                                detalle += f" Captura: {ruta_captura}"
                            except Exception:
                                logging.exception(
                                    "No se pudo guardar la captura de factura %s",
                                    factura
                                )

                            logging.error(
                                "Factura %s falló tras los reintentos.",
                                factura
                            )

                except ValueError as error:
                    estado = "RECHAZADA"
                    detalle = str(error)
                    logging.warning(
                        "Fila %s rechazada por validación: %s",
                        numero_fila, detalle
                    )

                except Exception as error:
                    estado = "FALLA_SISTEMA"
                    detalle = str(error)
                    logging.exception(
                        "Error inesperado en la fila %s", numero_fila
                    )

                    nombre_captura = re.sub(
                        r"[^a-zA-Z0-9_-]", "_", str(factura)
                    )
                    try:
                        page.screenshot(
                            path=f"salida/evidencias/factura_{nombre_captura}.png",
                            full_page=True
                        )
                    except Exception:
                        logging.exception("No se pudo guardar la captura.")

                try:
                    id_lista_out = int(id_lista) if id_lista is not None else ""
                except ValueError:
                    id_lista_out = id_lista

                try:
                    nit_out = int(limpiar_nit(nit)) if limpiar_nit(nit) else ""
                except ValueError:
                    nit_out = limpiar_nit(nit)

                hoja_salida.append([
                    id_lista_out,
                    factura,
                    nit_out,
                    estado,
                    codigo_causacion,
                    detalle
                ])
                logging.info(
                    "Resultado factura %s: %s. %s",
                    factura, estado, detalle
                )

                # Guardar después de cada factura para no perder avances.
                guardar_resultados(wb_salida, hoja_salida)

            print(f"\nProceso terminado. Resultados: {ruta_salida}")
            print(f"Registro de ejecución: {RUTA_LOG}")

        finally:
            browser.close()
            wb_entrada.close()


if __name__ == "__main__":
    preparar_y_procesar()